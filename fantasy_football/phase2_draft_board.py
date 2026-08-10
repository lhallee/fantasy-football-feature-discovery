"""Build draft-board pages from the generated 2026 player catalog."""

from __future__ import annotations

import argparse
import json

from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .players_2026 import PLAYERS, PlayerProjection


DRAFT_POSITIONS = ("QB", "RB", "WR", "TE", "K")
UNAVAILABLE_STATUSES = frozenset({"CUT", "RET"})
SHEET_ORDER = ("Overall", *DRAFT_POSITIONS)
SHEET_COLORS = {
    "Overall": "17365D",
    "QB": "1F4E78",
    "RB": "548235",
    "WR": "C65911",
    "TE": "7030A0",
    "K": "BF9000",
}
WORKBOOK_COLUMNS = (
    ("Draft Status", "drafted", 15),
    ("Fantasy Team", "fantasy_team", 18),
    ("Draft Round", "draft_round", 12),
    ("Draft Pick", "draft_pick", 11),
    ("Notes", "notes", 28),
    ("Player", "name", 24),
    ("NFL Team", "team", 10),
    ("Pos", "position", 8),
    ("Predicted Pts", "predicted_points", 14),
    ("Low 80%", "prediction_low_80", 12),
    ("High 80%", "prediction_high_80", 12),
    ("Previous Pts", "previous_season_points", 14),
    ("Pos Rank", "position_rank", 10),
    ("Overall Rank", "overall_point_rank", 12),
    ("Team Pos Rank", "team_position_rank", 13),
    ("Depth Tier", "depth_tier", 11),
    ("Depth Rank", "depth_rank", 11),
    ("Roster Status", "roster_status", 14),
    ("Confidence", "confidence", 24),
    ("OOD", "out_of_distribution", 9),
    ("Age", "age", 9),
    ("Experience", "years_experience", 11),
    ("NFL Draft Year", "draft_year", 14),
    ("NFL Draft Round", "nfl_draft_round", 15),
    ("NFL Draft Pick", "nfl_draft_pick", 13),
    ("College", "college", 24),
    ("Player ID", "player_id", 18),
)


def _draft_record(player: PlayerProjection) -> dict[str, Any]:
    """Return one flat, spreadsheet-ready player record."""
    return {
        "drafted": "Available",
        "fantasy_team": "",
        "draft_round": None,
        "draft_pick": None,
        "notes": "",
        "name": player.name,
        "team": player.team,
        "position": player.model_position,
        "predicted_points": player.predicted_points,
        "prediction_low_80": player.prediction_interval_80_low,
        "prediction_high_80": player.prediction_interval_80_high,
        "previous_season_points": player.previous_season_points,
        "position_rank": player.position_rank,
        "overall_point_rank": player.overall_point_rank,
        "team_position_rank": player.team_position_rank,
        "depth_tier": player.depth_tier,
        "depth_rank": player.depth_rank,
        "roster_status": player.roster_status,
        "confidence": player.confidence,
        "out_of_distribution": player.out_of_distribution,
        "age": player.age,
        "years_experience": player.years_experience,
        "draft_year": player.draft_year,
        "nfl_draft_round": player.draft_round,
        "nfl_draft_pick": player.draft_number,
        "college": player.college,
        "player_id": player.player_id,
    }


def _predicted_order(player: PlayerProjection) -> tuple[float, str, str]:
    """Return the deterministic descending-prediction sort key."""
    predicted_points = player.predicted_points
    if predicted_points is None:
        raise ValueError(f"Draft-eligible player {player.player_id} has no prediction.")
    return (-predicted_points, player.name.casefold(), player.player_id)


def draft_board_pages(
    players: Iterable[PlayerProjection] = PLAYERS,
    *,
    available_only: bool = True,
) -> dict[str, tuple[dict[str, Any], ...]]:
    """Return overall and position pages sorted by predicted points."""
    eligible = tuple(  # (n_eligible_players,)
        player
        for player in players
        if player.fantasy_eligible
        and player.model_position in DRAFT_POSITIONS
        and player.predicted_points is not None
        and (
            not available_only
            or (player.roster_status or "").upper() not in UNAVAILABLE_STATUSES
        )
    )
    ordered = tuple(sorted(eligible, key=_predicted_order))  # (n_eligible_players,)
    pages: dict[str, tuple[dict[str, Any], ...]] = {
        "Overall": tuple(_draft_record(player) for player in ordered)
    }
    for position in DRAFT_POSITIONS:
        position_players = (  # (n_position_players,)
            player for player in ordered if player.model_position == position
        )
        pages[position] = tuple(_draft_record(player) for player in position_players)
    return pages


def write_draft_board_json(
    output_path: Path,
    players: Iterable[PlayerProjection] = PLAYERS,
    *,
    available_only: bool = True,
) -> Path:
    """Write the workbook-neutral draft-board payload."""
    pages = draft_board_pages(players, available_only=available_only)
    payload: Mapping[str, Any] = {
        "season": 2026,
        "sort_contract": "predicted_points_descending",
        "available_only": available_only,
        "pages": pages,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return output_path


def _style_workbook_sheet(sheet: Any, sheet_name: str, row_count: int) -> None:
    """Apply the verified draft-board presentation and interaction layer."""
    accent = SHEET_COLORS[sheet_name]
    last_row = row_count + 4
    last_column = len(WORKBOOK_COLUMNS)
    last_letter = sheet.cell(row=4, column=last_column).column_letter

    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = accent
    sheet.freeze_panes = "F5"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet["A1"] = f"2026 Fantasy Football Draft Board | {sheet_name}"
    sheet["A1"].fill = PatternFill("solid", fgColor=accent)
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet["A2"] = (
        "ESPN full-PPR preseason projections | Forecast origin 2026-08-09 | "
        "Sorted only by predicted season points"
    )
    sheet["A2"].fill = PatternFill("solid", fgColor="D9EAF7")
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color="1F1F1F")
    sheet["A2"].comment = Comment(
        "Generated directly from fantasy_football.players_2026.PLAYERS.",
        "Fantasy Football Discovery",
    )
    sheet.row_dimensions[2].height = 22

    sheet.merge_cells("A3:D3")
    sheet["A3"] = f'=SUBTOTAL(103,$F$5:$F${last_row})&" visible players"'
    sheet.merge_cells("E3:H3")
    sheet["E3"] = "Editable draft controls: columns A:E"
    sheet.merge_cells(f"I3:{last_letter}3")
    caution = "Approximate 80% intervals; confidence and OOD are model diagnostics"
    if sheet_name == "K":
        caution += "; kicker evidence is mixed"
    sheet["I3"] = caution
    for cell in (sheet["A3"], sheet["E3"], sheet["I3"]):
        cell.fill = PatternFill("solid", fgColor="EAF2F8")
        cell.font = Font(name="Aptos", size=10, bold=True, color="17365D")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[3].height = 23

    for column_index, (label, _, width) in enumerate(WORKBOOK_COLUMNS, start=1):
        cell = sheet.cell(row=4, column=column_index, value=label)
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.font = Font(name="Aptos", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="FFFFFF"))
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.row_dimensions[4].height = 36

    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    for row in sheet.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.font = Font(name="Aptos", size=10)
            cell.alignment = Alignment(vertical="center")
        for cell in row[:5]:
            cell.fill = editable_fill
    for column in ("I", "J", "K", "L", "U"):
        for (cell,) in sheet[f"{column}5:{column}{last_row}"]:
            cell.number_format = "0.0"

    status = DataValidation(
        type="list",
        formula1='"Available,Watch,My Team,Taken"',
        allow_blank=True,
    )
    sheet.add_data_validation(status)
    status.add(f"A5:A{last_row}")
    draft_round = DataValidation(
        type="whole", operator="between", formula1="1", formula2="30", allow_blank=True
    )
    sheet.add_data_validation(draft_round)
    draft_round.add(f"C5:C{last_row}")
    draft_pick = DataValidation(
        type="whole", operator="between", formula1="1", formula2="500", allow_blank=True
    )
    sheet.add_data_validation(draft_pick)
    draft_pick.add(f"D5:D{last_row}")

    editable_range = f"A5:E{last_row}"
    sheet.conditional_formatting.add(
        editable_range,
        FormulaRule(
            formula=['$A5="My Team"'],
            fill=PatternFill("solid", fgColor="C6EFCE"),
            font=Font(color="006100", bold=True),
        ),
    )
    sheet.conditional_formatting.add(
        editable_range,
        FormulaRule(
            formula=['$A5="Taken"'],
            fill=PatternFill("solid", fgColor="D9D9D9"),
            font=Font(color="666666", strike=True),
        ),
    )
    sheet.conditional_formatting.add(
        f"I5:I{last_row}",
        DataBarRule(start_type="min", end_type="max", color=accent, showValue=True),
    )
    sheet["A4"].comment = Comment(
        "Track Available, Watch, My Team, or Taken with the dropdown.",
        "Fantasy Football Discovery",
    )
    sheet["I4"].comment = Comment(
        "Predicted full-season ESPN full-PPR points. This is the only sort key.",
        "Fantasy Football Discovery",
    )

    table = Table(displayName=f"DraftBoard{sheet_name}", ref=f"A4:{last_letter}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.auto_filter.ref = table.ref
    sheet.print_area = f"A1:{last_letter}{last_row}"
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def verify_draft_workbook(
    workbook_path: Path,
    pages: Mapping[str, tuple[dict[str, Any], ...]],
) -> None:
    """Reopen a workbook and verify sheet, order, value, and control contracts."""
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    if tuple(workbook.sheetnames) != SHEET_ORDER:
        raise AssertionError(f"Unexpected draft-board sheets: {workbook.sheetnames!r}")
    for sheet_name in SHEET_ORDER:
        sheet = workbook[sheet_name]
        expected = pages[sheet_name]  # (n_sheet_players,)
        scores = [sheet.cell(row=row, column=9).value for row in range(5, sheet.max_row + 1)]
        identifiers = [
            sheet.cell(row=row, column=len(WORKBOOK_COLUMNS)).value
            for row in range(5, sheet.max_row + 1)
        ]
        if scores != [record["predicted_points"] for record in expected]:
            raise AssertionError(f"{sheet_name} workbook predictions changed.")
        if scores != sorted(scores, reverse=True):
            raise AssertionError(f"{sheet_name} workbook is not prediction-sorted.")
        if identifiers != [record["player_id"] for record in expected]:
            raise AssertionError(f"{sheet_name} workbook player order changed.")
        if sheet.freeze_panes != "F5" or len(sheet.tables) != 1:
            raise AssertionError(f"{sheet_name} workbook controls are incomplete.")
        if len(sheet.data_validations.dataValidation) != 3:
            raise AssertionError(f"{sheet_name} workbook validation is incomplete.")
    workbook.close()


def write_draft_workbook(
    output_path: Path,
    players: Iterable[PlayerProjection] = PLAYERS,
    *,
    available_only: bool = True,
) -> Path:
    """Write and independently verify the organized multi-sheet draft workbook."""
    pages = draft_board_pages(players, available_only=available_only)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "2026 Fantasy Football Draft Board"
    workbook.properties.subject = "ESPN full-PPR projections and draft tracker"
    workbook.properties.creator = "Fantasy Football Raw-Feature Discovery"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    for sheet_name in SHEET_ORDER:
        sheet = workbook.create_sheet(sheet_name)
        rows = pages[sheet_name]  # (n_sheet_players,)
        for row_index, record in enumerate(rows, start=5):
            for column_index, (_, key, _) in enumerate(WORKBOOK_COLUMNS, start=1):
                sheet.cell(row=row_index, column=column_index, value=record[key])
        _style_workbook_sheet(sheet, sheet_name, len(rows))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    verify_draft_workbook(output_path, pages)
    return output_path


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    """Parse the standalone draft-board output command."""
    parser = argparse.ArgumentParser(
        description="Build the 2026 fantasy draft board as JSON or Excel."
    )
    parser.add_argument("--format", choices=("json", "xlsx"), default="xlsx")
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    """Write the requested draft-board artifact and print its path."""
    arguments = _parse_args(argv)
    if arguments.output is None:
        output_path = (
            Path("experiments/phase2/artifacts/draft_board_2026.json")
            if arguments.format == "json"
            else Path("outputs/fantasy_football_draft_board_2026.xlsx")
        )
    else:
        output_path = arguments.output
    if arguments.format == "json":
        output_path = write_draft_board_json(output_path)
    else:
        output_path = write_draft_workbook(output_path)
    print(output_path.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
