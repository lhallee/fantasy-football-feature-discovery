"""Validate the completed Phase 3 scientific and draft artifacts."""

import hashlib
import json
import joblib
import numpy as np
import pandas as pd
import pytest

from pathlib import Path

from openpyxl import load_workbook
from sklearn.metrics import (
    average_precision_score,
    brier_score_loss,
    log_loss,
    roc_auc_score,
)

from fantasy_football.phase3_injury import build_injury_feature_bundle
from fantasy_football.phase3_runner import _verify_prefit_inputs


ROOT = Path(__file__).parents[1]
PHASE3 = ROOT / "experiments" / "phase3"
RESULT = json.loads((PHASE3 / "result_summary.json").read_text(encoding="utf-8"))
RUN = ROOT / RESULT["run_directory"]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_phase3_output_hashes_match_every_run_artifact() -> None:
    declared = json.loads((RUN / "output_hashes.json").read_text(encoding="utf-8"))
    actual = {
        path.relative_to(RUN).as_posix(): _sha256(path)
        for path in sorted(RUN.rglob("*"))
        if path.is_file() and path.name != "output_hashes.json"
    }

    assert actual == declared


def test_current_injury_overlay_hashes_and_counts_match_release() -> None:
    release = json.loads(
        (PHASE3 / "experiment_manifest_v1.json").read_text(encoding="utf-8")
    )
    overlay = json.loads(
        (PHASE3 / "current_injury_manifest_v1.json").read_text(encoding="utf-8")
    )

    assert _sha256(PHASE3 / "result_summary.json") == release["result_summary"]["sha256"]
    assert _sha256(ROOT / release["workbook"]["path"]) == release["workbook"]["sha256"]
    assert (
        _sha256(ROOT / release["current_injury_screen"]["manifest_path"])
        == release["current_injury_screen"]["manifest_sha256"]
    )
    assert overlay["available_players_screened"] == 951
    assert overlay["current_injury_listings"] == 109
    assert overlay["players_without_current_listing"] == 842
    for record in overlay["files"].values():
        path = ROOT / record["path"]
        assert path.stat().st_size == record["bytes"]
        assert _sha256(path) == record["sha256"]


def test_prefit_inputs_and_saved_model_reproduce_current_risks() -> None:
    prefit = json.loads(
        (PHASE3 / "prefit_manifest_v1.json").read_text(encoding="utf-8")
    )
    injury_directory = PHASE3 / "data" / "raw"
    operational_manifest = json.loads(
        (PHASE3 / "current_injury_manifest_v1.json").read_text(encoding="utf-8")
    )
    current_draft_hash = _sha256(ROOT / "fantasy_football" / "phase3_draft.py")
    overlaid_prefit = json.loads(json.dumps(prefit))
    overlaid_prefit["code_sha256"]["fantasy_football/phase3_draft.py"] = (
        current_draft_hash
    )
    _verify_prefit_inputs(ROOT, injury_directory, overlaid_prefit)

    assert (
        prefit["code_sha256"]["fantasy_football/phase3_draft.py"]
        != current_draft_hash
    )
    for relative, expected in operational_manifest["operational_code_sha256"].items():
        assert _sha256(ROOT / relative) == expected

    bundle = build_injury_feature_bundle(ROOT, injury_directory)
    current_mask = bundle.metadata["target_season"].eq(2026)  # (14631,)
    current_X = bundle.frame.loc[current_mask]  # (958, 500)
    current_keys = bundle.metadata.loc[
        current_mask,
        ["target_season", "player_id", "model_position"],
    ].copy()  # (958, 3)
    model = joblib.load(RUN / "injury_model.joblib")
    replayed_probability = model.predict_proba(current_X)  # (958,)
    replayed = current_keys.assign(
        replayed_probability=replayed_probability
    )  # (958, 4)
    stored = pd.read_parquet(
        RUN / "injury_adjusted_players_2026.parquet",
        columns=[
            "target_season",
            "player_id",
            "model_position",
            "injury_probability",
        ],
    )  # (958, 4)
    comparison = stored.merge(
        replayed,
        on=["target_season", "player_id", "model_position"],
        how="left",
        validate="one_to_one",
    )  # (958, 5)

    assert len(model.feature_columns) == current_X.shape[1] == 500
    np.testing.assert_allclose(
        comparison["injury_probability"],
        comparison["replayed_probability"],
        atol=1e-15,
        rtol=0.0,
    )


def test_audit_metrics_recompute_from_player_predictions() -> None:
    predictions = pd.read_parquet(RUN / "audit_predictions.parquet")  # (3417, 9)
    y = predictions["physical_injury_reported"].to_numpy(dtype="int8")  # (3417,)
    p = predictions["injury_probability"].to_numpy(dtype="float64")  # (3417,)
    metric_p = np.clip(p, 1e-6, 1.0 - 1e-6)  # (3417,)
    metrics = RESULT["audit"]["selected_pooled"]

    assert roc_auc_score(y, metric_p) == pytest.approx(metrics["roc_auc"], abs=1e-12)
    assert average_precision_score(y, metric_p) == pytest.approx(
        metrics["average_precision"], abs=1e-12
    )
    assert log_loss(y, metric_p, labels=[0, 1]) == pytest.approx(
        metrics["log_loss"], abs=1e-12
    )
    assert brier_score_loss(y, metric_p) == pytest.approx(
        metrics["brier_score"], abs=1e-12
    )


def test_negative_control_and_exposure_claims_match_saved_rows() -> None:
    null = pd.read_csv(RUN / "negative_control_summary.csv")  # (11, 3)
    diagnostic = json.loads(
        (RUN / "posthoc_exposure_diagnostic.json").read_text(encoding="utf-8")
    )

    assert len(null) == 11
    assert 0.45 <= null["pooled_roc_auc"].mean() <= 0.55
    assert null["pooled_roc_auc"].mean() == pytest.approx(
        RESULT["negative_control_mean_pooled_roc_auc"], abs=1e-12
    )
    assert diagnostic["overall"]["selected_model_roc_auc"] == pytest.approx(
        RESULT["audit"]["selected_pooled"]["roc_auc"], abs=1e-12
    )
    assert diagnostic["precutoff_exposure_subsets"][-1]["rule"] == "lag1_games >= 8"


def test_current_combined_scores_follow_the_declared_formula() -> None:
    board = pd.read_parquet(
        RUN / "injury_adjusted_players_2026.parquet"
    )  # (958, 18)
    expected = 100.0 * (
        0.8 * board["fantasy_score_percentile"]
        + 0.2 * board["health_probability"]
    )  # (958,)

    assert len(board) == 958
    assert int(board["available_for_draft"].sum()) == 951
    np.testing.assert_allclose(board["combined_draft_score"], expected, atol=1e-12)
    assert board["injury_probability"].between(0.0, 1.0).all()
    assert board.groupby("model_position")["position_combined_rank"].min().eq(1).all()


def test_selected_lineage_has_no_current_outcome_source() -> None:
    lineage = json.loads((RUN / "feature_lineage.json").read_text(encoding="utf-8"))

    assert len(lineage) == RESULT["selected_feature_count"] == 500
    for feature, record in lineage.items():
        sources = set(record["sources"])
        offsets = record["source_offsets"]
        assert "modeling_table.target_roster" not in sources
        if feature.startswith("injury_"):
            assert min(offsets) >= 2
        for source, offset in zip(record["sources"], offsets, strict=True):
            if source in {"weekly_stats", "player_seasons"}:
                assert offset >= 1


def test_injury_adjusted_workbook_is_sorted_by_combined_score() -> None:
    path = ROOT / RESULT["workbook"]
    if not path.is_file():
        pytest.skip("Generated workbook is not present in this checkout.")
    workbook = load_workbook(path, read_only=False, data_only=True)

    assert workbook.sheetnames == ["ALL", "QB", "RB", "WR", "TE", "K", "INJURED"]
    expected_headers = (
        "Risk-Adjusted Draft Score",
        "Projected 2026 Fantasy Points",
        "Predicted Physical Injury-Report Probability",
    )
    for name in workbook.sheetnames:
        sheet = workbook[name]
        scores = [sheet.cell(row, 5).value for row in range(2, sheet.max_row + 1)]
        assert all(first >= second for first, second in zip(scores, scores[1:]))
        ranks = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
        assert ranks == list(range(1, len(ranks) + 1))
        assert tuple(sheet.cell(1, column).value for column in (5, 6, 7)) == expected_headers
        assert sheet.freeze_panes == "A2"
        assert sheet.column_dimensions["L"].hidden

    all_ids = {
        workbook["ALL"].cell(row, 12).value
        for row in range(2, workbook["ALL"].max_row + 1)
    }
    injured_ids = {
        workbook["INJURED"].cell(row, 12).value
        for row in range(2, workbook["INJURED"].max_row + 1)
    }
    source_board = pd.read_parquet(
        PHASE3 / "artifacts" / "injury_adjusted_players_2026.parquet"
    )  # (958, 18)
    available_ids = set(
        source_board.loc[source_board["available_for_draft"], "player_id"]
    )
    current_injuries = pd.read_csv(
        PHASE3 / "artifacts" / "current_injuries_2026-08-14.csv"
    )  # (n_injured, 20)

    assert all_ids.isdisjoint(injured_ids)
    assert all_ids | injured_ids == available_ids
    assert injured_ids == set(current_injuries["player_id"])
    for row in range(2, workbook["INJURED"].max_row + 1):
        assert workbook["INJURED"].cell(row, 16).value
        assert workbook["INJURED"].cell(row, 17).value == "2026-08-14"
        assert workbook["INJURED"].cell(row, 20).value
