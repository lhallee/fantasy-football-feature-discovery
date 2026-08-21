"""Test the injury-adjusted draft preference score and filters."""

import pandas as pd
import pytest

from openpyxl import load_workbook

from fantasy_football.phase3_draft import (
    SHEET_ORDER,
    combine_draft_scores,
    query_risk_adjusted_players,
    write_risk_adjusted_workbook,
)


def _predictions() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "player_id": ["a", "b", "c"],
            "candidate_name": ["Alpha", "Beta", "Charlie"],
            "team": ["BUF", "BUF", "NYJ"],
            "model_position": ["QB", "QB", "K"],
            "predicted_points": [300.0, 200.0, 150.0],
            "injury_probability": [0.50, 0.10, 0.20],
        }
    )  # (3, 6)


def test_combined_score_uses_declared_eighty_twenty_formula() -> None:
    board = combine_draft_scores(_predictions())  # (3, 13)
    alpha = board.loc[board["player_id"].eq("a")].iloc[0]
    beta = board.loc[board["player_id"].eq("b")].iloc[0]

    assert alpha["fantasy_score_percentile"] == 1.0
    assert beta["fantasy_score_percentile"] == 0.5
    assert alpha["combined_draft_score"] == pytest.approx(90.0)
    assert beta["combined_draft_score"] == pytest.approx(58.0)


def test_query_normalizes_pk_and_rejects_negative_limits() -> None:
    board = combine_draft_scores(_predictions())  # (3, 13)

    kickers = query_risk_adjusted_players(board, position="PK", limit=None)

    assert kickers["player_id"].tolist() == ["c"]
    with pytest.raises(ValueError, match="nonnegative"):
        query_risk_adjusted_players(board, limit=-1)


def test_workbook_partitions_current_injuries_without_read_me(tmp_path) -> None:
    board = combine_draft_scores(_predictions())  # (3, 13)
    current_injuries = pd.DataFrame(
        {
            "player_id": ["a"],
            "current_injury_or_designation": ["Hamstring"],
            "current_injury_status": ["Questionable"],
            "expected_return": ["Aug 22, 2026"],
            "evidence_date": ["2026-08-14"],
            "status_as_of": ["2026-08-14"],
            "source_publisher": ["Example feed"],
            "source_tier": ["Test evidence"],
            "source_url": ["https://example.com/injuries"],
            "source_player_url": [""],
        }
    )  # (1, 10)
    output_path = tmp_path / "draft.xlsx"

    write_risk_adjusted_workbook(
        board,
        output_path,
        current_injuries=current_injuries,
        injury_as_of="2026-08-14",
    )

    workbook = load_workbook(output_path, read_only=False, data_only=True)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    assert "Read Me" not in workbook.sheetnames
    assert workbook["ALL"].max_row == 3
    assert workbook["INJURED"].max_row == 2
    assert workbook["ALL"]["E1"].value == "Risk-Adjusted Draft Score"
    assert workbook["ALL"]["F1"].value == "Projected 2026 Fantasy Points"
    assert (
        workbook["ALL"]["G1"].value
        == "Predicted Physical Injury-Report Probability"
    )
    assert workbook["ALL"]["L2"].value != "a"
    assert workbook["INJURED"]["L2"].value == "a"
    assert workbook["INJURED"]["Q2"].value == "2026-08-14"
    assert workbook["ALL"].column_dimensions["L"].hidden
